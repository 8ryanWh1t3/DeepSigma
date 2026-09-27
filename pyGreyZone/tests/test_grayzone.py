"""Boundary checks for the local v0.1 assessment and governance path."""

import importlib.util
import json
import tempfile
import unittest
from dataclasses import replace
from datetime import datetime, timedelta, timezone
from pathlib import Path

from deepsigma.grayzone import (AssessmentConfig, FieldMap, GrayZoneEpisode,
                                approve_patch, assess, cerpa_packet, compare,
                                exposed_dependencies, load_events, load_lattice_jsonl,
                                normalize, propose_patch, record_apply,
                                review_hypothesis)
from deepsigma.grayzone.bridge import core_claim, core_event, core_review
from deepsigma.grayzone.cli import main
from deepsigma.grayzone.graph import graph_json, ntriples
from deepsigma.grayzone.report import markdown, to_dict, write_excel


EXAMPLE = Path(__file__).resolve().parents[1] / "examples" / "synthetic_cuas.lattice.jsonl"


class AssessmentTests(unittest.TestCase):
    def test_pattern_is_unreviewed_and_location_alone_is_insufficient(self):
        events = load_lattice_jsonl(EXAMPLE)
        result = assess(events)
        self.assertEqual(result, assess(reversed(events)))
        self.assertEqual(len(result.events), 4)
        self.assertEqual(len(result.hypotheses), 1)
        candidate = result.hypotheses[0]
        self.assertEqual(candidate.event_ids, ("obs-001", "obs-002", "obs-003"))
        self.assertEqual(candidate.status, "unreviewed")
        self.assertEqual(len(candidate.alternatives), 3)
        self.assertEqual(len(candidate.source_groups), 3)
        self.assertTrue(all("obs-004" not in (l.left, l.right) for l in result.links))
        self.assertTrue(all(len(e.sources[0].sha256) == 64 for e in result.events))

    def test_lineage_gate_and_unlinked_weak_signal(self):
        events = load_lattice_jsonl(EXAMPLE)[:3]
        copies = [replace(e, sources=(replace(e.sources[0], lineage_group="same-upstream"),)) for e in events]
        result = assess(copies)
        self.assertFalse(result.hypotheses)
        self.assertTrue(result.weak_signals)

    def test_window_identity_and_time_zone(self):
        events = load_lattice_jsonl(EXAMPLE)
        far = replace(events[0], id="later", occurred_at=datetime(2026, 12, 1, tzinfo=timezone.utc))
        result = assess([*events, far], AssessmentConfig(window=timedelta(days=1)))
        self.assertEqual([e.id for e in result.events], ["later"])
        self.assertFalse(result.hypotheses)
        self.assertNotEqual(result.id, assess(events).id)
        with self.assertRaisesRegex(ValueError, "offset"):
            normalize({"id": "x", "occurred_at": "2026-09-18T13:00:00",
                       "kind": "x", "channel": "x"})

    def test_broad_context_and_duplicate_id(self):
        e1, e2 = load_lattice_jsonl(EXAMPLE)[:2]
        e2 = replace(e2, entities=(), tags=(), assets=())
        self.assertFalse(assess([e1, e2]).links)
        with self.assertRaisesRegex(ValueError, "duplicate event id"):
            assess([e1, replace(e1, summary="changed")])
        conflicting_source = replace(e1.sources[0], sha256="0" * 64)
        with self.assertRaisesRegex(ValueError, "source record changed content"):
            assess([e1, replace(e1, id="other", sources=(conflicting_source,))])

    def test_csv_and_error_location(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "events.csv"
            source.write_text("id,occurred_at,kind,channel,source_id,entities,tags\n"
                              "a,2026-09-18T13:00:00Z,report,human,observer,track-1|track-2,odd|repeat\n")
            event = load_events(source)[0]
            self.assertEqual(event.entities, ("track-1", "track-2"))
            self.assertEqual(event.tags, ("odd", "repeat"))
            bad = Path(temp) / "bad.jsonl"
            bad.write_text('{"id":"a"}\n')
            with self.assertRaisesRegex(ValueError, "bad.jsonl:1"):
                load_events(bad)
            bad.write_text("not-json\n")
            with self.assertRaisesRegex(ValueError, "bad.jsonl:1"):
                load_events(bad)
            with self.assertRaisesRegex(ValueError, "unknown field map"):
                FieldMap.from_mapping({"bogus": "field"})

    def test_cerpa_requires_review_and_approval(self):
        result = GrayZoneEpisode(load_lattice_jsonl(EXAMPLE)).assess()
        claim = result.hypotheses[0]
        review = review_hypothesis(result, claim.id, reviewer="analyst", verdict="inconclusive",
                                   rationale="Check ordinary activity", reviewed_at="2026-09-18T14:00:00Z")
        patch = propose_patch(review, proposer="analyst", action="collect",
                              description="Request original records")
        packet = cerpa_packet(result, review, patch)
        self.assertIsNone(packet["apply"])
        self.assertEqual(packet["patch"]["status"], "proposed")
        self.assertIsInstance(json.dumps(packet), str)
        with self.assertRaisesRegex(ValueError, "approval"):
            record_apply(patch, approver="lead", applied_at="2026-09-18T15:00:00Z",
                         outcome="recorded", note="Collected")
        approved = approve_patch(patch, approver="lead", approved_at="2026-09-18T14:30:00Z")
        applied = record_apply(approved, approver="lead", applied_at="2026-09-18T15:00:00Z",
                               outcome="recorded", note="Original records received")
        self.assertEqual(cerpa_packet(result, review, approved, applied)["apply"]["patch_id"], patch.id)
        self.assertEqual(core_event(result.events[0])["domain"], "intelops")
        self.assertIs(core_claim(claim, result)["metadata"]["coordination_established"], False)
        self.assertEqual(core_review(review)["verdict"], "inconclusive")

    def test_reports_graph_drift_cli(self):
        events = load_lattice_jsonl(EXAMPLE)
        result = assess(events)
        self.assertEqual(to_dict(result)["evidence_index"]["obs-001"][0]["record_id"], "a-001")
        self.assertIn("coordination remains unproven", markdown(result))
        self.assertTrue(graph_json(result)["edges"])
        self.assertIn("supportedBy", ntriples(result))
        self.assertEqual(compare(assess(events[:2]), result)["new_events"], ["obs-003", "obs-004"])
        self.assertEqual(exposed_dependencies("A", {"A": ["B"], "B": ["C"], "C": ["A"]}),
                         {"B": ("A", "B"), "C": ("A", "B", "C")})
        with tempfile.TemporaryDirectory() as temp:
            paths = [Path(temp) / name for name in ("report.json", "report.md", "graph.json", "graph.nt")]
            self.assertEqual(main([str(EXAMPLE), "--lattice", "--json", str(paths[0]),
                                   "--markdown", str(paths[1]), "--graph", str(paths[2]),
                                   "--rdf", str(paths[3])]), 0)
            self.assertTrue(all(path.stat().st_size for path in paths))
            self.assertEqual(len(json.loads(paths[0].read_text())["hypotheses"]), 1)

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "optional openpyxl not installed")
    def test_excel_tabs_and_formula_escape(self):
        import openpyxl
        events = load_lattice_jsonl(EXAMPLE)
        result = assess([replace(events[0], summary='=HYPERLINK("bad")'), *events[1:]])
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "cerpa.xlsx"
            write_excel(result, path)
            book = openpyxl.load_workbook(path)
            self.assertEqual(book.sheetnames, ["Dashboard", "C", "E", "R", "P", "A", "Memory"])
            self.assertTrue(book["E"]["H2"].value.startswith("'=HYPERLINK"))
            self.assertEqual(book["P"].max_row, 1)
            self.assertEqual(book["A"].max_row, 1)


if __name__ == "__main__":
    unittest.main()

"""Decision-ready JSON, Markdown and optional CERPA Excel export."""

from __future__ import annotations

import json
from dataclasses import asdict, is_dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from .provenance import evidence_index
from .schema import Assessment


def _default(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat().replace("+00:00", "Z")
    if is_dataclass(value):
        return asdict(value)
    raise TypeError(f"cannot encode {type(value).__name__}")


def to_dict(assessment: Assessment) -> dict[str, Any]:
    result = json.loads(json.dumps(assessment, default=_default, ensure_ascii=False))
    result["evidence_index"] = evidence_index(assessment.events)
    return result


def write_json(assessment: Assessment, path: str | Path) -> None:
    Path(path).write_text(json.dumps(to_dict(assessment), indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def _cell(value: object) -> str:
    return str(value).replace("|", "\\|").replace("\n", " ")


def markdown(assessment: Assessment) -> str:
    lines = ["# Gray Zone Pattern Assessment", "", f"Assessment: `{assessment.id}`",
             f"Window: {assessment.window_start} to {assessment.window_end}",
             f"Observations: {len(assessment.events)} · Links: {len(assessment.links)} · Candidates: {len(assessment.hypotheses)}",
             "", "## Candidate patterns", ""]
    if not assessment.hypotheses:
        lines += ["No candidate passed the configured minimum gates.", ""]
    for h in assessment.hypotheses:
        lines += [f"### {h.id} · priority {h.priority_score}/100 · {h.status}", "",
                  h.statement, "", "Evidence event IDs: " + ", ".join(f"`{i}`" for i in h.event_ids),
                  "Source lineage groups: " + ", ".join(h.source_groups), "",
                  "Competing explanations:"]
        for alt in h.alternatives:
            lines.append(f"- {_cell(alt.explanation)} — Check: {_cell(alt.discriminating_check)}")
        lines += ["", "Collection gaps:"]
        lines += ["- " + _cell(gap) for gap in h.collection_gaps]
        lines.append("")
    lines += ["## Observations and provenance", "", "| Event | Time (UTC) | Channel | Source record / SHA-256 | Summary |",
              "|---|---|---|---|---|"]
    for e in assessment.events:
        source = "; ".join(f"{s.source_id}/{s.record_id} @ {s.locator} ({s.sha256})" for s in e.sources)
        lines.append("| " + " | ".join(_cell(v) for v in (e.id, e.occurred_at.isoformat(), e.channel, source, e.summary)) + " |")
    lines += ["", "## Limits", ""] + ["- " + item for item in assessment.limitations]
    return "\n".join(lines) + "\n"


def write_excel(assessment: Assessment, path: str | Path) -> None:
    """Seven-tab Excel reduction: Dashboard, C, E, R, P, A, Memory."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("Excel export requires pip install 'deepsigma-grayzone[excel]'") from exc

    def safe(value: object) -> object:
        if isinstance(value, datetime):
            value = value.isoformat()
        if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
            return "'" + value
        return value

    wb = Workbook()
    wb.remove(wb.active)
    sheets: dict[str, list[list[object]]] = {
        "Dashboard": [["Assessment", assessment.id], ["Method", assessment.method],
                      ["Observations", len(assessment.events)], ["Candidate patterns", len(assessment.hypotheses)],
                      ["Caution", "Priority score is review order, not probability or proof of coordination."]],
        "C": [["Claim ID", "Status", "Priority", "Statement", "Event IDs", "Source groups"]]
             + [[h.id, h.status, h.priority_score, h.statement, ", ".join(h.event_ids),
                 ", ".join(h.source_groups)] for h in assessment.hypotheses],
        "E": [["Event ID", "UTC", "Kind", "Channel", "Entity IDs", "Asset IDs", "Location", "Summary", "Record hashes"]]
             + [[e.id, e.occurred_at, e.kind, e.channel, ", ".join(e.entities), ", ".join(e.assets),
                 e.location, e.summary, "; ".join(s.sha256 for s in e.sources)] for e in assessment.events],
        "R": [["Claim ID", "Reviewer", "Verdict", "Rationale", "Alternative checks", "Collection gaps"]]
             + [[h.id, "", "unreviewed", "", "; ".join(a.discriminating_check for a in h.alternatives),
                 "; ".join(h.collection_gaps)] for h in assessment.hypotheses],
        "P": [["Patch ID", "Review ID", "Proposer", "Action", "Status", "Description"]],
        "A": [["Apply ID", "Patch ID", "Approver", "Time", "Outcome", "Note"]],
        "Memory": [["Event ID", "Source", "Record ID", "Lineage group", "SHA-256", "Locator"]]
                  + [[e.id, s.source_id, s.record_id, s.independent_group, s.sha256, s.locator]
                     for e in assessment.events for s in e.sources],
    }
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append([safe(v) for v in row])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="394836")
        for col in ws.columns:
            width = min(58, max(12, max(len(str(c.value or "")) for c in col) + 2))
            ws.column_dimensions[get_column_letter(col[0].column)].width = width
    wb.save(path)
